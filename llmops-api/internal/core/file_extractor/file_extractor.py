import csv
import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from charset_normalizer import from_bytes
from injector import inject
from langchain_core.documents import Document as LCDocument

from internal.exception import FailException
from internal.model import UploadFile
from internal.service.cos_service import CosService


@inject
@dataclass
class FileExtractor:
  """提取文件正文；使用各格式解析器，避免 Unstructured 的旧依赖冲突。"""

  cos_service: CosService

  def load(self, upload_file: UploadFile, return_text=False, is_unstructured=True):
    with tempfile.TemporaryDirectory() as directory:
      path = Path(directory) / Path(upload_file.key).name
      self.cos_service.download_file(upload_file.key, str(path))
      return self.load_from_file(str(path), return_text, is_unstructured)

  @classmethod
  def load_from_url(cls, url: str, return_text=False):
    response = requests.get(url, timeout=30)
    response.raise_for_status()
    with tempfile.TemporaryDirectory() as directory:
      path = Path(directory) / (Path(urlparse(url).path).name or 'download.txt')
      path.write_bytes(response.content)
      return cls.load_from_file(str(path), return_text)

  @classmethod
  def load_from_file(cls, file_path: str, return_text=False, is_unstructured=True):
    path = Path(file_path)
    extension = path.suffix.lower()
    texts = []
    if extension == '.pdf':
      from pypdf import PdfReader

      texts = [page.extract_text() or '' for page in PdfReader(path).pages]
    elif extension == '.docx':
      from docx import Document

      document = Document(path)
      texts = [paragraph.text for paragraph in document.paragraphs]
      texts.extend(
        '\t'.join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
      )
    elif extension in {'.doc', '.ppt'}:
      executable = shutil.which('soffice') or shutil.which('libreoffice')
      if not executable:
        raise FailException('解析 .doc/.ppt 需要安装 LibreOffice，或先转成 .docx/.pptx')
      with tempfile.TemporaryDirectory() as directory:
        output_format = 'docx' if extension == '.doc' else 'pptx'
        subprocess.run(
          [
            executable,
            f'-env:UserInstallation={Path(directory).as_uri()}/profile',
            '--headless',
            '--convert-to',
            output_format,
            '--outdir',
            directory,
            str(path.resolve()),
          ],
          check=True,
          timeout=60,
          capture_output=True,
        )
        return cls.load_from_file(
          str(Path(directory) / f'{path.stem}.{output_format}'), return_text
        )
    elif extension == '.pptx':
      from pptx import Presentation

      texts = [
        '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        for slide in Presentation(path).slides
      ]
    elif extension == '.xlsx':
      from openpyxl import load_workbook

      workbook = load_workbook(path, read_only=True, data_only=True)
      try:
        texts = [
          '\n'.join(
            '\t'.join('' if value is None else str(value) for value in row)
            for row in sheet.iter_rows(values_only=True)
          )
          for sheet in workbook
        ]
      finally:
        workbook.close()
    elif extension == '.xls':
      import xlrd

      workbook = xlrd.open_workbook(str(path))
      try:
        texts = [
          '\n'.join(
            '\t'.join(map(str, sheet.row_values(index))) for index in range(sheet.nrows)
          )
          for sheet in workbook.sheets()
        ]
      finally:
        workbook.release_resources()
    else:
      decoded = from_bytes(path.read_bytes()).best()
      if decoded is None:
        raise FailException('无法识别文件编码')
      text = str(decoded)
      if extension in {'.html', '.htm', '.xml'}:
        soup = BeautifulSoup(text, 'html.parser')
        for element in soup(['script', 'style']):
          element.decompose()
        text = soup.get_text('\n', strip=True)
      elif extension == '.csv':
        text = '\n'.join('\t'.join(row) for row in csv.reader(io.StringIO(text)))
      texts = [text]
    documents = [
      LCDocument(page_content=text, metadata={'source': str(path)})
      for text in texts
      if text.strip()
    ]
    return (
      '\n\n'.join(document.page_content for document in documents)
      if return_text
      else documents
    )
