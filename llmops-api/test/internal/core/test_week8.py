import json
import uuid
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from langchain_core.language_models.fake_chat_models import FakeListChatModel
from langchain_core.messages import AIMessageChunk, HumanMessage
from langchain_core.tools import tool

from config import Config
from internal.core.agent.agents import AgentQueueManager, FunctionCallAgent
from internal.core.agent.entities.agent_entity import AgentConfig
from internal.core.agent.entities.queue_entity import QueueEvent
from internal.core.file_extractor import FileExtractor
from internal.entity.conversation_entity import InvokeFrom
from internal.entity.dataset_entity import DEFAULT_PROCESS_RULE
from internal.model import Account
from internal.schema.document_schema import CreateDocumentsReq
from internal.service import AppService, JwtService, ProcessRuleService


def queue_manager():
  redis = Mock()
  redis.get.return_value = None
  return AgentQueueManager(uuid.uuid4(), uuid.uuid4(), InvokeFrom.DEBUGGER, redis)


def test_agent_stream_and_long_term_memory():
  manager = queue_manager()
  agent = FunctionCallAgent(
    AgentConfig(
      llm=FakeListChatModel(responses=['你好']), enable_long_term_memory=True
    ),
    manager,
  )
  events = list(agent.run('问候', [], '用户喜欢中文'))
  assert events[0].event == QueueEvent.LONG_TERM_MEMORY_RECALL
  assert events[0].observation == '用户喜欢中文'
  assert ''.join(event.answer for event in events) == '你好'
  assert events[-1].messages[0]['type'] == 'system'


def test_agent_background_exception_terminates_stream():
  agent = FunctionCallAgent(
    AgentConfig(llm=FakeListChatModel(responses=['unused'])), queue_manager()
  )
  events = list(agent.run('你好', [HumanMessage('不成对的历史')]))
  assert len(events) == 1
  assert events[0].event == QueueEvent.ERROR


def test_agent_tool_call_round_trip():
  @tool
  def add_one(number: int) -> int:
    """Add one to a number."""
    return number + 1

  class ToolModel(FakeListChatModel):
    def bind_tools(self, tools):
      return self

    def stream(self, messages, **kwargs):
      if messages[-1].type == 'tool':
        yield AIMessageChunk(content=messages[-1].content)
      else:
        yield AIMessageChunk(
          content='',
          tool_call_chunks=[
            {'name': 'add_one', 'args': '{"number": 2}', 'id': 'call-1', 'index': 0}
          ],
        )

  agent = FunctionCallAgent(
    AgentConfig(llm=ToolModel(responses=[]), tools=[add_one]), queue_manager()
  )
  events = list(agent.run('2 加 1'))
  assert [event.event for event in events] == [
    QueueEvent.AGENT_THOUGHT,
    QueueEvent.AGENT_ACTION,
    QueueEvent.AGENT_MESSAGE,
  ]
  assert events[1].tool_input == {'number': 2}
  assert events[-1].answer == '3'


def test_agent_input_and_output_review():
  input_review_agent = FunctionCallAgent(
    AgentConfig(
      llm=FakeListChatModel(responses=['不应调用']),
      review_config={
        'enable': True,
        'keywords': ['敏感词'],
        'inputs_config': {'enable': True, 'preset_response': '内容无法处理'},
        'outputs_config': {'enable': False},
      },
    ),
    queue_manager(),
  )
  events = list(input_review_agent.run('这里有敏感词'))
  assert [event.answer for event in events] == ['内容无法处理']

  output_review_agent = FunctionCallAgent(
    AgentConfig(
      llm=FakeListChatModel(responses=['包含秘密的回答']),
      review_config={
        'enable': True,
        'keywords': ['秘密'],
        'inputs_config': {'enable': False, 'preset_response': ''},
        'outputs_config': {'enable': True},
      },
    ),
    queue_manager(),
  )
  assert ''.join(event.answer for event in output_review_agent.run('正常问题')) == (
    '包含**的回答'
  )


def test_queue_timeout_and_stop():
  manager = queue_manager()
  assert list(manager.listen(timeout=0.01))[-1].event == QueueEvent.TIMEOUT
  manager = queue_manager()
  manager.redis_client.get.return_value = b'1'
  assert list(manager.listen(timeout=2))[-1].event == QueueEvent.STOP


def test_agent_sse_endpoint(client, db, monkeypatch):
  account_id = uuid.uuid4()
  db.session.add(
    Account(
      id=account_id,
      name='测试账号',
      email='agent@example.com',
      avatar='',
      last_login_ip='127.0.0.1',
    )
  )
  db.session.flush()
  monkeypatch.setenv('JWT_SECRET_KEY', 'test-jwt-secret-at-least-32-bytes')
  token = JwtService.generate_token({'sub': str(account_id)})
  monkeypatch.setattr(
    AppService,
    'debug_chat',
    lambda self, app_id, query, account: iter(
      ['event: agent_message\ndata: {"answer": "你好"}\n\n']
    ),
  )
  response = client.post(
    f'/apps/{uuid.uuid4()}/debug',
    json={'query': '你好'},
    headers={'Authorization': f'Bearer {token}'},
  )
  assert response.mimetype == 'text/event-stream'
  events = [
    json.loads(line[6:])
    for line in response.get_data(as_text=True).splitlines()
    if line.startswith('data: ')
  ]
  assert ''.join(event['answer'] for event in events) == '你好'


@pytest.mark.parametrize(
  'extension, content, expected',
  [
    ('txt', '知识库测试', '知识库测试'),
    ('md', '# 标题\n知识库', '知识库'),
    ('csv', 'name,value\n知识库,1', '知识库\t1'),
    ('html', '<script>bad()</script><p>知识库</p>', '知识库'),
  ],
)
def test_extract_text_formats(tmp_path, extension, content, expected):
  path = tmp_path / f'test.{extension}'
  path.write_text(content)
  result = FileExtractor.load_from_file(str(path), return_text=True)
  assert expected in result
  assert 'bad()' not in result


def test_extract_office_and_pdf(tmp_path):
  from docx import Document
  from openpyxl import Workbook
  from pptx import Presentation
  from pypdf import PdfWriter

  document = Document()
  document.add_paragraph('知识库文档')
  document.save(tmp_path / 'test.docx')
  workbook = Workbook()
  workbook.active.append(['知识库表格', 1])
  workbook.save(tmp_path / 'test.xlsx')
  deck = Presentation()
  slide = deck.slides.add_slide(deck.slide_layouts[0])
  slide.shapes.title.text = '知识库幻灯片'
  deck.save(tmp_path / 'test.pptx')
  pdf = PdfWriter()
  pdf.add_blank_page(width=100, height=100)
  pdf.write(tmp_path / 'test.pdf')
  for extension in ['docx', 'xlsx', 'pptx']:
    assert '知识库' in FileExtractor.load_from_file(
      str(tmp_path / f'test.{extension}'), True
    )
  assert FileExtractor.load_from_file(str(tmp_path / 'test.pdf'), True) == ''


def test_process_rule_validation_and_splitting(app):
  with app.test_request_context(
    method='POST',
    json={'upload_file_ids': [str(uuid.uuid4())], 'process_type': 'automatic'},
  ):
    req = CreateDocumentsReq()
    assert req.validate(), req.errors
    assert req.rule.data == DEFAULT_PROCESS_RULE['rule']
  rule = SimpleNamespace(rule=DEFAULT_PROCESS_RULE['rule'])
  cleaned = ProcessRuleService.clean_text_by_process_rule(
    '你好  世界 https://example.com a@b.com', rule
  )
  assert cleaned.strip() == '你好 世界'
  splitter = ProcessRuleService.get_text_splitter_by_process_rule(rule)
  assert len(splitter.split_text('知识库。' * 600)) > 1
  with app.test_request_context(
    method='POST',
    json={
      'upload_file_ids': [str(uuid.uuid4())],
      'process_type': 'custom',
      'rule': {'pre_process_rules': [None]},
    },
  ):
    req = CreateDocumentsReq()
    assert not req.validate()


def test_celery_redis_auth_and_ssl(monkeypatch):
  monkeypatch.setenv('REDIS_USE_SSL', 'true')
  monkeypatch.setenv('REDIS_USERNAME', 'user@name')
  monkeypatch.setenv('REDIS_PASSWORD', 'pass:/word')
  config = Config()
  assert config.CELERY['broker_url'].startswith('rediss://user%40name:pass%3A%2Fword@')
  assert config.CELERY['broker_url'].endswith('?ssl_cert_reqs=required')
